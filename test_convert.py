"""ทดสอบแปลง Word/Excel + ดึงตาราง/รูปภาพ"""
import sys, tempfile
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parent))

import fitz
from PIL import Image
from PySide6.QtWidgets import QApplication
from app.main import MainWindow
from app.pdfops import PdfDocument, find_thai_font

tmp = Path(tempfile.mkdtemp())

# เอกสารทดสอบ: หน้า 1 มีตารางจริง (มีเส้น) + ข้อความไทย, หน้า 2 มีรูปฝังอยู่
doc = fitz.open()
page = doc.new_page(width=595, height=842)
page.insert_font(fontname="thai", fontfile=find_thai_font())
page.insert_text((60, 60), "ใบเสนอราคา", fontname="thai", fontsize=16)

rows, cols = 4, 3
x0, y0, w, h = 60, 100, 150, 30
data = [
    ["รายการ", "จำนวน", "ราคา"],
    ["Notebook Dell", "1", "38,200.00"],
    ["Mouse", "2", "590.00"],
    ["รวม", "3", "38,790.00"],
]
for r in range(rows):
    for c in range(cols):
        rect = fitz.Rect(x0 + c * w, y0 + r * h, x0 + (c + 1) * w, y0 + (r + 1) * h)
        page.draw_rect(rect, color=(0, 0, 0), width=0.8)
        page.insert_textbox(rect + (4, 6, -4, -4), data[r][c],
                            fontname="thai", fontsize=10)

page2 = doc.new_page(width=595, height=842)
logo = tmp / "logo.png"
im = Image.new("RGB", (120, 60), (200, 30, 30))
im.save(logo)
page2.insert_image(fitz.Rect(60, 60, 180, 120), filename=str(logo))
page2.insert_text((60, 200), "page two", fontsize=12)

src = tmp / "quote.pdf"
doc.save(str(src)); doc.close()

app = QApplication(sys.argv)


def fresh():
    w = MainWindow()
    w.doc = PdfDocument(src)
    w.viewer.set_document(w.doc)
    w._refresh_all(keep_row=0)
    return w


# ---- หาตาราง ----
d = PdfDocument(src)
tables = d.find_tables(0)
assert len(tables) == 1, f"ควรเจอ 1 ตาราง ได้ {len(tables)}"
assert len(tables[0]) == 4, tables[0]
flat = " ".join(" ".join(r) for r in tables[0])
assert "Notebook Dell" in flat, flat
assert "รายการ" in flat, "หัวตารางภาษาไทยต้องอ่านออก"
assert "38,200.00" in flat, flat
print("OK  หาตารางเจอ อ่านค่าถูก (รวมภาษาไทย)")

assert d.find_tables(1) == [], "หน้าที่ไม่มีตารางต้องคืนลิสต์ว่าง"
print("OK  หน้าที่ไม่มีตารางไม่หลอนว่าเจอ")

# ---- Excel ----
xlsx = tmp / "out.xlsx"
count = d.to_excel(xlsx)
assert count == 1 and xlsx.exists()
from openpyxl import load_workbook
book = load_workbook(xlsx)
sheet = book[book.sheetnames[0]]
values = [[c.value for c in row] for row in sheet.iter_rows()]
assert values[0][0] == "รายการ", values[0]
assert values[1][2] == "38,200.00", values[1]
assert sheet.max_row == 4 and sheet.max_column == 3
print(f"OK  Excel: {book.sheetnames} {sheet.max_row}x{sheet.max_column}")

# หน้าที่ไม่มีตาราง -> ต้องเตือน ไม่ใช่สร้างไฟล์เปล่า
try:
    d.to_excel(tmp / "empty.xlsx", indexes=[1])
except ValueError as exc:
    assert "ไม่พบตาราง" in str(exc)
    assert not (tmp / "empty.xlsx").exists(), "ไม่เจอตารางต้องไม่สร้างไฟล์ทิ้งไว้"
    print("OK  ไม่เจอตารางแล้วเตือน ไม่สร้างไฟล์เปล่า")
else:
    raise AssertionError("ควร error")

# ---- Word ----
docx = tmp / "out.docx"
d.to_word(docx)
assert docx.exists() and docx.stat().st_size > 0
from docx import Document
text = "\n".join(p.text for p in Document(docx).paragraphs)
tables_docx = Document(docx).tables
assert tables_docx, "ต้องแปลงตารางมาเป็นตารางใน Word ด้วย"
cells = " ".join(c.text for t in tables_docx for r in t.rows for c in r.cells)
assert "Notebook Dell" in cells, cells[:200]
assert "รายการ" in cells, "ข้อความไทยต้องแปลงมาถูก"
print("OK  Word: มีตารางและข้อความไทยครบ")

# แปลงเฉพาะบางหน้า
docx2 = tmp / "page2.docx"
d.to_word(docx2, indexes=[1])
body = "\n".join(p.text for p in Document(docx2).paragraphs)
assert "page two" in body, body
assert "Notebook Dell" not in body, "เลือกหน้า 2 ต้องไม่ติดเนื้อหาหน้า 1"
print("OK  Word เฉพาะหน้าที่เลือก")

# ---- ลายเซ็นที่ยังไม่บันทึก ต้องติดไปใน Word ด้วย ----
d2 = PdfDocument(src)
sig = tmp / "sig.png"
sp = fitz.open(); spp = sp.new_page(width=200, height=80)
spp.draw_bezier((10, 60), (60, 5), (120, 90), (190, 25), color=(0, 0, 1), width=3)
spp.get_pixmap(alpha=False).save(str(sig)); sp.close()
d2.add_stamp(0, str(sig), fitz.Rect(300, 600, 500, 680))
docx3 = tmp / "signed.docx"
d2.to_word(docx3)
from docx import Document as D3
assert len(D3(docx3).inline_shapes) >= 1, "ลายเซ็นต้องติดไปใน Word ด้วย"
assert len(d2.stamps) == 1, "แปลงแล้วต้องไม่ทำให้ลายเซ็นในโปรแกรมหาย"
print("OK  ลายเซ็นที่ยังไม่บันทึกติดไปใน Word ด้วย")

# ---- ดึงรูปที่ฝังในไฟล์ ----
folder = tmp / "imgs"
written = d.extract_images([0, 1], folder)
assert len(written) == 1, [p.name for p in written]
assert written[0].name.startswith("หน้า2"), written[0].name
with Image.open(written[0]) as got:
    assert got.size == (120, 60), got.size
print(f"OK  ดึงรูปในไฟล์: {written[0].name} {got.size}")

assert d.extract_images([0], tmp / "none") == [], "หน้าที่ไม่มีรูปต้องได้ลิสต์ว่าง"
print("OK  หน้าที่ไม่มีรูปคืนลิสต์ว่าง")

# ---- ผ่าน UI ----
w = fresh()
out_x = tmp / "ui.xlsx"
with patch("app.main.QInputDialog.getItem", return_value=("ทุกหน้า (2 หน้า)", True)), \
     patch("app.main.QFileDialog.getSaveFileName", return_value=(str(out_x), "")):
    w.export_excel()
assert out_x.exists()
print("OK  ดึงตารางผ่านเมนูได้")

w = fresh()
w.thumbs.setCurrentRow(1)              # ไปหน้า 2 (ไม่มีตาราง)
with patch("app.main.QInputDialog.getItem", return_value=("หน้าปัจจุบัน (หน้า 2)", True)), \
     patch("app.main.QFileDialog.getSaveFileName", return_value=(str(tmp / "nope.xlsx"), "")), \
     patch("app.main.QMessageBox.information") as info:
    w.export_excel()
    assert info.called, "ไม่เจอตารางต้องแจ้งผู้ใช้"
print("OK  ไม่เจอตารางแล้วแจ้งผู้ใช้ผ่าน UI")

w = fresh()
with patch("app.main.QInputDialog.getItem", return_value=("ทุกหน้า (2 หน้า)", True)), \
     patch("app.main.QFileDialog.getExistingDirectory", return_value=str(tmp / "uiimgs")):
    w.export_embedded_images()
assert list((tmp / "uiimgs").glob("*.png"))
print("OK  ดึงรูปผ่านเมนูได้")

print("\nผ่านทั้งหมด")
